"""对问题二最终策略、因果场景索引和 result2.xlsx 做独立交叉校验。"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import compare_q2_methods as compare
import q2_causal_core as core


ROOT = Path(__file__).resolve().parents[2]
FINAL_BOOK = ROOT / "result" / "2" / "result2.xlsx"
FINAL_STRATEGY = ROOT / "result" / "2" / "tmp" / "q2_final_solution.npz"
COMPARISON = ROOT / "result" / "2" / "tmp" / "q2_method_comparison.csv"
AUDIT_PATH = ROOT / "result" / "2" / "tmp" / "q2_final_audit.json"


def main() -> None:
    data = core.load_inputs()
    strategy = core.load_strategy(FINAL_STRATEGY)
    metrics = core.evaluate_strategy(data, strategy)
    indices = core.output_indices(data)

    charge = strategy.charge_kwh[indices]
    discharge = strategy.discharge_kwh[indices]
    soc = strategy.soc_end_kwh[indices]
    previous_soc = np.column_stack(
        (strategy.soc_start_kwh[indices], soc[:, :-1])
    )
    soc_residual = (
        soc
        - previous_soc
        - core.BATTERY_CHARGE_EFFICIENCY * charge
        + discharge / core.BATTERY_DISCHARGE_EFFICIENCY
    )
    max_soc_residual = float(np.max(np.abs(soc_residual)))
    max_combined_power_kw = float(np.max(charge + discharge) / core.DT_HOURS)

    if metrics["supply_interruptions"] != 0:
        raise AssertionError("最终策略存在供电中断。")
    if metrics["simultaneous_charge_discharge_intervals"] != 0:
        raise AssertionError("最终策略存在同一时段同时充放电。")
    if max_soc_residual > 1.0e-6 or max_combined_power_kw > 5000.0 + 1.0e-6:
        raise AssertionError("储能状态转移或功率边界不满足。")

    forecasts = compare.load_forecasts()
    if forecasts is None:
        raise AssertionError("严格因果预测缓存缺失。")
    ridge = forecasts["ridge_causal"]
    max_future_index_violation = -len(data.dates)
    for day in indices:
        _, _, _, residual_days = core.make_residual_scenarios(data, ridge, int(day), 6)
        _, _, _, similar_days = core.make_direct_similar_scenarios(data, int(day), 6)
        max_future_index_violation = max(
            max_future_index_violation,
            int(np.max(residual_days - day)),
            int(np.max(similar_days - day)),
        )
    if max_future_index_violation >= 0:
        raise AssertionError("场景构造使用了目标日或未来日期。")

    comparison = pd.read_csv(COMPARISON)
    eligible = comparison.loc[comparison["参与可实施方法排名"]]
    eligible_zero = eligible.loc[eligible["供电中断次数"] == 0]
    selected = eligible_zero.sort_values("总购电费（元）").iloc[0]
    if selected["方法编号"] != strategy.method_id:
        raise AssertionError("最终策略不是零中断候选中的最低费用方法。")

    workbook = load_workbook(FINAL_BOOK, read_only=True, data_only=True)
    if workbook.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise AssertionError("最终工作簿的工作表结构改变。")
    plan_sheet = workbook["计划购电量"]
    workbook_plan = np.array(
        [
            [float(value or 0.0) for value in row]
            for row in plan_sheet.iter_rows(
                min_row=2, max_row=335, min_col=2, max_col=145, values_only=True
            )
        ]
    )
    max_workbook_plan_error = float(
        np.max(np.abs(workbook_plan - strategy.planned_purchase_kwh[indices]))
    )
    emergency_sheet = workbook["紧急购电量"]
    workbook_emergency_sum = float(
        sum(float(row[0] or 0.0) for row in emergency_sheet.iter_rows(min_row=2, min_col=3, max_col=3, values_only=True))
    )
    expected_emergency_sum = float(strategy.emergency_kwh[indices].sum())
    if max_workbook_plan_error > 1.0e-9:
        raise AssertionError("工作簿计划购电量与最终策略不一致。")
    if abs(workbook_emergency_sum - expected_emergency_sum) > 1.0e-6:
        raise AssertionError("工作簿紧急购电合计与最终策略不一致。")

    audit = {
        "模型版本": core.MODEL_VERSION,
        "入选方法": strategy.name,
        "候选方法数（可实施）": int(len(eligible)),
        "全部可实施方法供电中断次数": [int(value) for value in eligible["供电中断次数"]],
        "场景日期相对目标日的最大索引差": int(max_future_index_violation),
        "最大供需平衡残差（kWh）": metrics["max_supply_balance_violation_kwh"],
        "最大SOC状态转移残差（kWh）": max_soc_residual,
        "最大充放电合计功率（kW）": max_combined_power_kw,
        "SOC最小值（kWh）": metrics["soc_min_kwh"],
        "SOC最大值（kWh）": metrics["soc_max_kwh"],
        "日末SOC最大误差（kWh）": metrics["max_daily_terminal_soc_error_kwh"],
        "同段充放电时段数": metrics["simultaneous_charge_discharge_intervals"],
        "工作簿计划购电最大写入误差（kWh）": max_workbook_plan_error,
        "工作簿紧急购电合计误差（kWh）": abs(workbook_emergency_sum - expected_emergency_sum),
        "供电中断次数": metrics["supply_interruptions"],
        "校验通过": True,
    }
    AUDIT_PATH.write_text(
        json.dumps(core.json_ready(audit), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(core.json_ready(audit), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
