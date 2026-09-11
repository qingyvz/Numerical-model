"""把问题二严格因果最优策略写入附件 5 的 result2.xlsx 模板并校验。"""

from __future__ import annotations

import copy
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "CUMCM2026Problems" / "C题" / "附件" / "附件5" / "result2.xlsx"
PAYLOAD = ROOT / "result" / "2" / "tmp" / "q2_workbook_payload.json"
OUTPUT = ROOT / "result" / "2" / "result2.xlsx"
QA_JSON = ROOT / "result" / "2" / "tmp" / "result2_validation.json"


def copy_cell_style(source, target) -> None:
    target._style = copy.copy(source._style)
    target.number_format = source.number_format
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)


def snapshot_row_styles(sheet: Worksheet, rows: list[int]) -> list[list[Any]]:
    return [[copy.copy(sheet.cell(row, column)._style) for column in range(1, sheet.max_column + 1)] for row in rows]


def apply_row_style(sheet: Worksheet, row: int, styles: list[Any], height: float | None) -> None:
    for column, style in enumerate(styles, start=1):
        sheet.cell(row, column)._style = copy.copy(style)
    if height is not None:
        sheet.row_dimensions[row].height = height


def parse_date(text: str) -> datetime:
    return datetime.strptime(text, "%Y-%m-%d")


def fill_planned_purchase(sheet: Worksheet, rows: list[dict[str, Any]]) -> None:
    if sheet.max_row != 335 or sheet.max_column != 147:
        raise ValueError(f"计划购电量模板规模异常：{sheet.max_row}×{sheet.max_column}")
    if len(rows) != 334:
        raise ValueError(f"计划购电输出日数应为334，实际为{len(rows)}")
    for offset, item in enumerate(rows, start=2):
        values = item["interval_kwh"]
        if len(values) != 144:
            raise ValueError(f"{item['date']} 的计划购电量不是144个时段")
        sheet.cell(offset, 1, parse_date(item["date"]))
        sheet.cell(offset, 1).number_format = "yyyy/m/d"
        for column, value in enumerate(values, start=2):
            sheet.cell(offset, column, float(value))
            sheet.cell(offset, column).number_format = "0.00"
        sheet.cell(offset, 146, float(item["daily_purchase_kwh"]))
        sheet.cell(offset, 147, float(item["daily_scheduled_cost_yuan"]))
        sheet.cell(offset, 146).number_format = "0.00"
        sheet.cell(offset, 147).number_format = "0.00"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:EQ{sheet.max_row}"


def fill_storage(sheet: Worksheet, rows: list[dict[str, Any]]) -> None:
    if len(rows) != 334 * 6:
        raise ValueError(f"充放电量行数应为2004，实际为{len(rows)}")
    style_rows = list(range(2, 8))
    styles = snapshot_row_styles(sheet, style_rows)
    heights = [sheet.row_dimensions[row].height for row in style_rows]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for index, item in enumerate(rows):
        row = index + 2
        pattern = index % 6
        apply_row_style(sheet, row, styles[pattern], heights[pattern])
        sheet.cell(row, 1, parse_date(item["date"]) if item["date"] else None)
        sheet.cell(row, 2, item["interval"])
        sheet.cell(row, 3, float(item["charge_kwh"]))
        sheet.cell(row, 4, float(item["discharge_kwh"]))
        sheet.cell(row, 5, item["soc_time"])
        sheet.cell(row, 6, float(item["soc_kwh"]) if item["soc_kwh"] is not None else None)
        if item["date"]:
            sheet.cell(row, 1).number_format = "yyyy/m/d"
        for column in (3, 4, 6):
            sheet.cell(row, column).number_format = "0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"


def fill_emergency(sheet: Worksheet, rows: list[dict[str, Any]]) -> None:
    if len(rows) < 334:
        raise ValueError(f"紧急购电记录至少应覆盖334天，实际仅{len(rows)}行")
    style = snapshot_row_styles(sheet, [2])[0]
    height = sheet.row_dimensions[2].height
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for index, item in enumerate(rows, start=2):
        apply_row_style(sheet, index, style, height)
        sheet.cell(index, 1, parse_date(item["date"]) if item["date"] else None)
        sheet.cell(index, 2, item["interval"])
        sheet.cell(index, 3, float(item["purchase_kwh"]))
        if item["date"]:
            sheet.cell(index, 1).number_format = "yyyy/m/d"
        sheet.cell(index, 3).number_format = "0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:C{sheet.max_row}"


def validate_workbook(path: Path, payload: dict[str, Any]) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    expected_sheets = ["计划购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(f"工作表名称异常：{workbook.sheetnames}")

    plan = workbook["计划购电量"]
    storage = workbook["充放电量"]
    emergency = workbook["紧急购电量"]
    if (plan.max_row, plan.max_column) != (335, 147):
        raise ValueError("计划购电量工作表规模不正确")
    if (storage.max_row, storage.max_column) != (2005, 6):
        raise ValueError("充放电量工作表规模不正确")
    expected_emergency_rows = len(payload["emergency_rows"])
    if (emergency.max_row, emergency.max_column) != (expected_emergency_rows + 1, 3):
        raise ValueError("紧急购电量工作表规模不正确")

    max_plan_total_error = 0.0
    max_plan_cost_error = 0.0
    for row_index, item in enumerate(payload["planned_purchase_rows"], start=2):
        interval_values = [float(plan.cell(row_index, column).value) for column in range(2, 146)]
        max_plan_total_error = max(
            max_plan_total_error,
            abs(sum(interval_values) - float(plan.cell(row_index, 146).value)),
        )
        max_plan_cost_error = max(
            max_plan_cost_error,
            abs(float(plan.cell(row_index, 147).value) - float(item["daily_scheduled_cost_yuan"])),
        )

    date_values = [plan.cell(row, 1).value.date().isoformat() for row in range(2, 336)]
    if date_values[0] != "2025-02-01" or date_values[-1] != "2025-12-31":
        raise ValueError("计划购电量日期范围不正确")
    if len(set(date_values)) != 334:
        raise ValueError("计划购电量日期存在重复")

    emergency_sum = sum(
        float(emergency.cell(row, 3).value or 0.0)
        for row in range(2, expected_emergency_rows + 2)
    )
    expected_emergency_sum = sum(
        float(item["purchase_kwh"]) for item in payload["emergency_rows"]
    )
    if abs(emergency_sum - expected_emergency_sum) > 1.0e-6:
        raise ValueError("紧急购电量写入合计与载荷不一致")
    expanded_dates: list[str] = []
    current_date: str | None = None
    for row in range(2, expected_emergency_rows + 2):
        value = emergency.cell(row, 1).value
        if value is not None:
            current_date = value.date().isoformat()
        if current_date is None:
            raise ValueError("紧急购电首行缺少日期")
        expanded_dates.append(current_date)
    if len(set(expanded_dates)) != 334 or expanded_dates[0] != "2025-02-01" or expanded_dates[-1] != "2025-12-31":
        raise ValueError("紧急购电记录未完整覆盖输出日期")

    formula_errors: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    formula_errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
    result = {
        "文件": str(path.relative_to(ROOT)),
        "工作表": expected_sheets,
        "计划购电日期数": 334,
        "计划购电时段数": 334 * 144,
        "充放电记录行数": 334 * 6,
        "紧急购电记录行数": expected_emergency_rows,
        "紧急购电量合计（kWh）": emergency_sum,
        "逐日计划购电量最大加总误差（kWh）": max_plan_total_error,
        "逐日计划购电费最大写入误差（元）": max_plan_cost_error,
        "公式错误": formula_errors,
        "校验通过": (
            max_plan_total_error <= 1.0e-6
            and max_plan_cost_error <= 1.0e-6
            and not formula_errors
        ),
    }
    if not result["校验通过"]:
        raise ValueError(f"工作簿校验失败：{result}")
    return result


def main() -> None:
    if not PAYLOAD.exists():
        raise FileNotFoundError("请先运行 program/2/solve_q2.py 生成写入载荷")
    payload = json.loads(PAYLOAD.read_text(encoding="utf-8"))
    workbook = load_workbook(TEMPLATE)
    if workbook.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise ValueError(f"模板工作表名称异常：{workbook.sheetnames}")
    fill_planned_purchase(workbook["计划购电量"], payload["planned_purchase_rows"])
    fill_storage(workbook["充放电量"], payload["storage_rows"])
    fill_emergency(workbook["紧急购电量"], payload["emergency_rows"])
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    validation = validate_workbook(OUTPUT, payload)
    QA_JSON.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
